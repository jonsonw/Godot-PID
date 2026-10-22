#!/usr/bin/env python3
# Copyright © 2026 Jonson Wang
# Build the four shipped Excel list templates. Run with the bundled Python that has openpyxl:
#   python3 tools/build_xlsx_templates.py
# 生成四套随包 Excel 清单模板。用带 openpyxl 的 Python 运行。
#
# Design / 设计：
#   - one worksheet per file, no cover sheet (product decision);
#   - rows 1-2 are the summary/title block with tokens, row 3 is the header, row 4 is the
#     single sample row, row 5 is the footer;
#   - A3 landscape, fit to one page wide, print titles rows 1:3, print area A1:last;
#   - headers are real labels (not tokens) so the file is readable when opened by a user;
#     the binder matches them by alias, and tokens are used in the summary/footer only.
# 每文件一张表，无封面；1-2 行是带令牌的汇总区，第 3 行表头，第 4 行唯一样例行，第 5 行页脚；
# A3 横向、按一页宽缩放、打印标题 1:3、打印区域 A1:last；表头是真实标签（用户打开可读），
# 绑定按别名识别，令牌只放汇总区/页脚。
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parents[1] / "assets" / "templates"

NAVY = "1F4E78"
HEADER = "2F5597"
LIGHT = "DDEBF7"
BORDER = "B7B7B7"

TITLE_FONT = Font(name="Arial", size=14, bold=True, color="FFFFFFFF")
SUB_FONT = Font(name="Arial", size=10, color="FF1F4E78")
HEAD_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
DATA_FONT = Font(name="Arial", size=9, color="FF000000")
FOOT_FONT = Font(name="Arial", size=9, color="FF404040")

THIN = Side(style="thin", color=BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Border(top=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)

# (key, header, width, numeric) / (字段, 表头, 宽, 是否数字)
LINE = [
    ("no", "序号", 6, True),
    ("line_no", "管线号", 16, False),
    ("line_kind", "管线类型", 10, False),
    ("from", "起点", 14, False),
    ("to", "终点", 14, False),
    ("from_equipment", "起点设备", 14, False),
    ("to_equipment", "终点设备", 14, False),
    ("dn", "公称尺寸", 10, False),
    ("spec", "管道等级", 10, False),
    ("medium", "介质", 12, False),
    ("insulation", "保温", 10, False),
    ("length_m", "长度(m)", 10, True),
    ("valves", "阀门", 18, False),
    ("valve_count", "阀门数", 8, True),
    ("segment_count", "管段数", 8, True),
    ("sheet", "图纸", 12, False),
]

EQUIPMENT = [
    ("no", "序号", 6, True),
    ("tag", "位号", 14, False),
    ("name", "名称", 20, False),
    ("type", "类型", 18, False),
    ("class", "标准类", 18, False),
    ("category", "类别", 12, False),
    ("qty", "数量", 8, True),
    ("material", "材质", 12, False),
    ("design_pressure", "设计压力", 12, True),
    ("design_temp", "设计温度", 10, True),
    ("volume", "容积", 10, True),
    ("nozzle_count", "管口数", 8, True),
    ("sheet", "图纸", 12, False),
    ("remark", "备注", 20, False),
]

INSTRUMENT = [
    ("no", "序号", 6, True),
    ("tag", "位号", 14, False),
    ("name", "名称", 18, False),
    ("type", "类型", 18, False),
    ("class", "标准类", 20, False),
    ("loop_no", "回路号", 10, False),
    ("measure_range", "量程", 14, False),
    ("signal_type", "信号类型", 10, False),
    ("accuracy", "精度", 8, False),
    ("power_supply", "电源", 12, False),
    ("connected_line", "连接管线", 14, False),
    ("connected_equipment", "连接设备", 24, False),
    ("sheet", "图纸", 12, False),
    ("remark", "备注", 20, False),
]

VALVE = [
    ("no", "序号", 6, True),
    ("tag", "位号", 14, False),
    ("name", "名称", 18, False),
    ("type", "类型", 16, False),
    ("class", "标准类", 18, False),
    ("line_no", "管线号", 14, False),
    ("nominal_size", "公称尺寸", 10, False),
    ("pressure_rating", "压力等级", 12, False),
    ("body_material", "阀体材质", 12, False),
    ("actuator_type", "执行机构", 12, False),
    ("fail_action", "故障动作", 12, False),
    ("qty", "数量", 8, True),
    ("sheet", "图纸", 12, False),
    ("remark", "备注", 18, False),
]

TEMPLATES = [
    ("line", LINE, "管线一览表", "G-PID 管线一览表"),
    ("equipment", EQUIPMENT, "设备一览表", "G-PID 设备一览表"),
    ("instrument", INSTRUMENT, "仪表索引表", "G-PID 仪表索引表"),
    ("valve", VALVE, "阀门一览表", "G-PID 阀门一览表"),
]


def build_one(path: Path, list_id: str, columns: list[tuple[str, str, int, bool]],
              sheet_name: str, title: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    last = get_column_letter(len(columns))

    # Summary / title block (tokens live here, never in the header row).
    # 汇总区/标题栏（令牌只在这里，不在表头行）。
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = LEFT
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = (
        "项目：{{project.name}}    图号：{{project.drawing_no}}    "
        f"版次：{{{{project.rev}}}}    共 {{{{list.{list_id}.count}}}} 条"
    )
    ws["A2"].font = SUB_FONT
    ws["A2"].fill = PatternFill("solid", fgColor=LIGHT)
    ws["A2"].alignment = LEFT
    ws.row_dimensions[2].height = 20

    # Header row 3. / 第 3 行表头。
    for idx, (key, header, width, numeric) in enumerate(columns, start=1):
        cell = ws.cell(row=3, column=idx, value=header)
        cell.font = HEAD_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.alignment = CENTER
        cell.border = BOX
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[3].height = 28

    # Single blank sample row 4: the filler clones this row and keeps its style.
    # 唯一空白样例行 4：填充器克隆它并保留样式。
    for idx, (key, header, width, numeric) in enumerate(columns, start=1):
        cell = ws.cell(row=4, column=idx)
        cell.value = None
        cell.font = DATA_FONT
        cell.alignment = RIGHT if numeric else LEFT
        cell.border = BOX
    ws.row_dimensions[4].height = 20
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{last}4"
    ws.sheet_view.showGridLines = False

    # Footer row 5 with merged blocks. / 第 5 行页脚，分块合并。
    footer_blocks = [
        ("编制：{{project.company}}", 1, min(3, len(columns))),
        ("审核：", 4, min(6, len(columns))),
        ("批准：", 7, min(9, len(columns))),
        ("日期：{{export.date}}", 10, len(columns)),
    ]
    for text, start, end in footer_blocks:
        if start > end:
            continue
        if start != end:
            ws.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        cell = ws.cell(row=5, column=start, value=text)
        cell.font = FOOT_FONT
        cell.alignment = LEFT
        cell.border = TOP
    for idx in range(1, len(columns) + 1):
        cell = ws.cell(row=5, column=idx)
        cell.border = TOP
    ws.row_dimensions[5].height = 22

    # Print setup: A3 landscape, one page wide, repeat the title/header rows.
    # 打印设置：A3 横向、一页宽、重复标题/表头行。
    ws.print_area = f"A1:{last}5"
    ws.print_title_rows = "1:3"
    ws.page_setup.paperSize = 8  # A3
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    wb.properties.creator = "G-PID"
    wb.properties.title = title
    wb.properties.subject = "G-PID 数据清单默认模板"
    wb.properties.description = "Tokenised summary/footer; header row 3; sample row 4; footer row 5."

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    for list_id, columns, sheet_name, title in TEMPLATES:
        path = OUT_DIR / f"gp_list_{list_id}_default.xlsx"
        build_one(path, list_id, columns, sheet_name, title)
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
